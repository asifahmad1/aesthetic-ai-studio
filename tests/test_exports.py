"""Export tests: PPTX / PDF / HTML validity and validation-dict structure.

These tests use only the local fixtures (no sibling modules required),
so they run even before the rest of the app is written.
"""
import html as html_lib
import os

import pytest
from pptx import Presentation
from pypdf import PdfReader

from export_pptx import export_pptx
from export_pdf import export_pdf
from export_html import export_html, slide_title


def _structure_ok(result):
    assert isinstance(result, dict), "export must return a dict"
    assert isinstance(result.get("ok"), bool), "'ok' must be a bool"
    checks = result.get("checks")
    assert isinstance(checks, list) and checks, "'checks' must be a non-empty list"
    for c in checks:
        assert set(c) >= {"name", "passed", "detail"}, \
            f"check missing keys: {c}"
        assert isinstance(c["passed"], bool)


def test_export_validation_dict_structure(project, theme, tmp_path):
    for fn, ext in ((export_pptx, "pptx"), (export_pdf, "pdf"),
                    (export_html, "html")):
        result = fn(project, theme, str(tmp_path / f"deck.{ext}"))
        _structure_ok(result)


def test_pptx_export_validates_and_opens(project, theme, tmp_path):
    out = str(tmp_path / "deck.pptx")
    result = export_pptx(project, theme, out)
    assert result["ok"], \
        "pptx validation failed: " + str([c for c in result["checks"]
                                          if not c["passed"]])
    assert os.path.getsize(out) > 10 * 1024
    prs = Presentation(out)  # must reopen cleanly
    assert len(prs.slides) == len(project["slides"])
    assert prs.slide_width > prs.slide_height  # 16:9 landscape


def test_pptx_uses_native_editable_text_not_screenshots(project, theme, tmp_path):
    out = str(tmp_path / "deck.pptx")
    export_pptx(project, theme, out)
    prs = Presentation(out)
    native_text_shapes = 0
    picture_shapes = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # picture
                picture_shapes += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                native_text_shapes += 1
    assert picture_shapes == 0, "export must never embed screenshots"
    assert native_text_shapes > 10, "deck should carry plenty of native text"
    assert any("Food Safety" in s.text_frame.text
               for s in prs.slides[0].shapes if s.has_text_frame)


def test_pptx_tables_and_bullets_and_notes(project, theme, tmp_path):
    out = str(tmp_path / "deck.pptx")
    export_pptx(project, theme, out)
    prs = Presentation(out)
    tables = sum(1 for s in prs.slides for sh in s.shapes if sh.has_table)
    assert tables >= 1, "table element must become a native pptx table"
    bullets = 0
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    if p.text.strip().startswith("•"):
                        bullets += 1
    assert bullets >= 5, "bullet lines must survive as bullet paragraphs"
    noted = sum(1 for s in prs.slides
                if s.notes_slide.notes_text_frame.text.strip())
    assert noted >= 3, "slide notes must land on notes slides"


def test_pptx_image_placeholders(project, theme, tmp_path):
    out = str(tmp_path / "deck.pptx")
    export_pptx(project, theme, out)
    prs = Presentation(out)
    found = any("image placeholder" in sh.text_frame.text.lower()
                for s in prs.slides for sh in s.shapes
                if sh.has_text_frame)
    assert found, "image elements must become labeled placeholders"


def test_pdf_export_page_count_and_header(project, theme, tmp_path):
    out = str(tmp_path / "deck.pdf")
    result = export_pdf(project, theme, out)
    assert result["ok"], \
        "pdf validation failed: " + str([c for c in result["checks"]
                                         if not c["passed"]])
    with open(out, "rb") as f:
        assert f.read(5) == b"%PDF-"
    reader = PdfReader(out)
    assert len(reader.pages) == len(project["slides"])
    w = float(reader.pages[0].mediabox.width)
    h = float(reader.pages[0].mediabox.height)
    assert w > h, "pages must be landscape"
    assert abs(w / h - 16 / 9) < 0.05, "pages must be 16:9"


def test_pdf_text_is_selectable(project, theme, tmp_path):
    out = str(tmp_path / "deck.pdf")
    export_pdf(project, theme, out)
    text = "\n".join(p.extract_text() or "" for p in PdfReader(out).pages)
    assert "Food Safety" in text, "PDF must contain real text, not images"
    assert "75" in text, "table content must appear in the PDF"


def test_html_export_contains_all_titles(project, theme, tmp_path):
    out = str(tmp_path / "deck.html")
    result = export_html(project, theme, out)
    assert result["ok"], \
        "html validation failed: " + str([c for c in result["checks"]
                                          if not c["passed"]])
    with open(out, encoding="utf-8") as f:
        content = f.read()
    assert "<section" in content
    for s in project["slides"]:
        title = slide_title(s)
        assert title in content or html_lib.escape(title) in content, \
            f"slide title missing from HTML: {title!r}"


def test_html_export_is_standalone(project, theme, tmp_path):
    out = str(tmp_path / "deck.html")
    export_html(project, theme, out)
    with open(out, encoding="utf-8") as f:
        content = f.read()
    assert "<!DOCTYPE html>" in content
    assert "<style>" in content, "CSS must be inline (standalone file)"
    assert 'rel="stylesheet"' not in content, "no external stylesheets"
    assert "<script src" not in content, "no external scripts"
    assert content.count("<section") >= len(project["slides"])


def test_exporters_never_claim_false_success(project, theme, tmp_path):
    # A project with zero slides still exports honestly; ok reflects checks.
    empty = {"id": "x", "title": "", "slides": []}
    r = export_pptx(empty, theme, str(tmp_path / "e.pptx"))
    _structure_ok(r)
    assert isinstance(r["ok"], bool)
    # Export where the parent path is blocked by a regular file must report
    # failure, not crash.
    blocker = tmp_path / "blocker"
    blocker.write_text("not a directory")
    bad = export_pdf(project, theme, str(blocker / "deck.pdf"))
    assert bad["ok"] is False
    assert any(not c["passed"] for c in bad["checks"])
